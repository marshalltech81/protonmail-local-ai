"""PPTX (PowerPoint .pptx) extractor.

Uses ``python-pptx`` to walk each slide in presentation order: the text
of every shape (text boxes, placeholders, auto shapes), each table one
line per row with its cells space-joined, and the shapes inside group
shapes, in the slide's shape order; then the slide's speaker notes.
Pictures are not read and nothing is OCR'd; charts and SmartArt keep
their text in parts this walk does not read. Each paragraph becomes one
line, separated by blank lines so the chunker has paragraph boundaries
to pack on, as in the DOCX extractor.

Slideshows (``.ppsx``), templates (``.potx``) and macro-enabled decks
(``.pptm``) are read the same way (#947), as are macro-enabled
slideshows (``.ppsm``) and templates (``.potm``) (#1042).
``pptx.Presentation`` accepts only the presentation and macro-enabled
presentation main-part types ("not a PowerPoint file" for the rest),
although python-pptx's own ``PartFactory`` already maps the slideshow
and template types too (``pptx/__init__.py``) to ``PresentationPart``.
So the package is opened with ``Package.open`` and its main part used
directly when it is one of the six presentation types. python-pptx has
no constant or mapping for the macro-enabled slideshow and template
main parts, so ``Package.open`` would load either as a generic part;
they are registered once, at import, in ``PartFactory.part_type_for``
(the extension point ``pptx/__init__.py`` itself uses), as the DOCX
extractor does for Word templates (#937). The payload bytes are not
changed. A macro-enabled file's ``vbaProject.bin`` is loaded by
python-pptx as an opaque part and never read: only slide and notes
text is. This route depends on python-pptx 1.0.2's ``Package`` /
``PartFactory`` API; ``test_python_pptx_route_for_variants_still_holds``
fails if an upgrade breaks it. Legacy binary ``.ppt`` is an OLE2
compound file, which the dispatcher records ``unsupported`` before this
module runs.

The work is bounded per extraction, counted as the walk goes:

* The package. The dispatcher's ZIP guard rejects a payload declaring
  more than ``ZIP_MAX_UNCOMPRESSED_BYTES`` before python-pptx opens it.
  Opening parses every XML part the package relates whole with lxml,
  which refuses elements nested more than 256 deep and text nodes over
  10 MB, and builds a part for every related member and walks every
  relationship. Plainly timed, 166 MB of slide XML parses in about a
  second but peaks at 2.5 GB, and 200,000 tiny related members take
  about 4 s. So before python-pptx opens it, a deck fails as
  ``PptxPackageBudgetError``, which the dispatcher records ``unsupported``
  since the same bytes always repeat it (#1032), when its members expand
  by more than
  ``_MAX_EXPANSION_BYTES`` past their compressed size, declare more than
  ``_MAX_DECLARED_BYTES`` together (#1033), number more than
  ``_MAX_MEMBERS``, or hold more than ``_MAX_RELS_BYTES`` of
  relationship parts, all read from the ZIP central directory. python-pptx
  follows the package's relationships recursively, so a long chain of
  related parts
  raises ``RecursionError`` while it opens, which is turned into
  ``PptxRelationshipChainError`` here: a chain of crafted parts is a
  property of the file, not host pressure.
* Slides. Each entry in the presentation's slide list costs one unit of
  ``_MAX_SLIDES``, and an entry naming a slide already read is skipped,
  so a list repeating one slide cannot multiply the walk.
* Shapes. Each shape costs one unit of ``_MAX_SHAPES``: a group and
  every shape inside it, at any depth, and each shape on a notes page.
  Groups are walked with an explicit stack, not recursion.
* Table cells. Each table row and each cell costs one unit of
  ``_MAX_TABLE_CELLS``.
* Text. Each character read costs one unit of ``_MAX_TEXT_CHARS``, and
  each paragraph and each element inside one costs ``_ELEMENT_COST``
  more, so empty runs and paragraphs cannot be read for free.

When a budget runs out the text collected so far is returned, as the
XLSX extractor does: the first budget to run out logs a rate-limited
WARNING naming it and is counted in the attachments aggregate's
``extractor_caps`` (#903).

The pre-open budgets trust the member sizes the ZIP central directory
declares, and a member that understates its size is still decompressed
whole when python-pptx reads it. So the whole extraction, budgets
included, runs in a child process under an address-space and a CPU
limit (``ooxml``, #1040): ``extract`` starts it, and ``extract_text``
is what runs in it. The child reports the budget that cut the text by
name, and this module logs it.
"""

from __future__ import annotations

import io
import logging
from collections.abc import Callable, Iterator

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.package import PartFactory
from pptx.oxml.ns import qn
from pptx.oxml.text import CT_RegularTextRun, CT_TextBody, CT_TextField, CT_TextLineBreak
from pptx.package import Package
from pptx.parts.presentation import PresentationPart
from pptx.presentation import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.slide import NotesSlide, Slide
from pptx.table import Table

from . import over_package_budget, warn_extractor_cap
from .ooxml import run_child

log = logging.getLogger("indexer.extractor.pptx")

# Address space (``RLIMIT_AS``), CPU seconds (``RLIMIT_CPU``) and
# wall-clock seconds the extraction's child process may use (#1040),
# past the CPU limit so a CPU-bound child meets that first. Plainly
# measured in the indexer image, child peak RSS and time:
#
# * a one-slide deck: 48 MB and 0.08 s, nearly all of it starting the
#   child and importing python-pptx (0.2 s before the image shipped
#   compiled bytecode, #1230; the cases below were measured then);
# * 1,000 slides with notes and 200 pictures (26 MB): 135 MB and 0.6 s;
# * 40 slides with 40 photos (30 MB): 106 MB and 0.3 s;
# * 150,000 empty text boxes on one slide (stopped at the shape
#   budget): 273 MB and 1.6 s;
# * 47 MiB of stored element-dense slide XML, inside the pre-open
#   budgets: 1,141 MB and 1.5 s;
# * a 0.5 MB deck whose slide declares a few KB but decompresses to
#   512 MiB (#1040): 1,071 MB and 0.6 s.
#
# 1 GiB is about 3.8 times the largest benign peak; the last two fail
# under it, as ``XMLSyntaxError`` (lxml reports a failed allocation as
# one) and ``MemoryError``. The CPU limit is many times the slowest
# case, as in the XLSX extractor.
CHILD_MAX_ADDRESS_SPACE_BYTES = 1024 * 1024 * 1024
CHILD_MAX_CPU_SECONDS = 30
CHILD_TIMEOUT_SECONDS = 45.0

# The walk budgets the child may report as having cut the text.
_CAP_NAMES = frozenset({"pptx_slides", "pptx_shapes", "pptx_table_cells", "pptx_text_chars"})

# Slide-list entries read. Far past any real deck; with repeated
# entries skipped, an entry costs microseconds.
_MAX_SLIDES = 5_000

# Shapes visited across slides and notes pages, group members included.
# Plainly timed, python-pptx builds a shape and reads its text in about
# 10 to 15 microseconds, so the budget stops a crafted slide in about
# 1.5 s; a 500-slide deck of 50 shapes a slide uses a quarter of it.
_MAX_SHAPES = 100_000

# Table rows and cells visited.
_MAX_TABLE_CELLS = 200_000

# Characters read, plus ``_ELEMENT_COST`` per paragraph and per element
# in one. Five times the dispatcher's default
# ``INDEXER_ATTACHMENT_MAX_EXTRACTED_CHARS`` (2,000,000), as in the XLSX
# extractor, so the dispatcher's cap still decides the stored length
# unless an operator raises it past this.
_MAX_TEXT_CHARS = 10_000_000

# What a paragraph or an element in one costs besides its characters.
# Plainly timed, reading a run costs about a microsecond however short
# its text, so a slide of empty runs stops after about 1,250,000 of them,
# in about 1.5 s. A run of ordinary text (tens of characters) still
# leaves several million characters of budget, past the dispatcher's cap.
_ELEMENT_COST = 8

# Bytes the package's members may expand by past their compressed size
# (#936, review round 1). python-pptx parses every XML part whole when it
# opens a deck: plainly timed, about 15 bytes of memory per byte of XML,
# so 166 MB of slide XML peaked at 2.5 GB. Pictures and media are stored
# already compressed and barely expand; a long text-heavy deck's XML
# expands by a few MB.
_MAX_EXPANSION_BYTES = 32 * 1024 * 1024

# Bytes all the package's members declare together, stored or compressed
# (#1033). A member stored uncompressed does not expand, so without this
# the payload's own stored XML (``INDEXER_ATTACHMENT_MAX_BYTES``) came on
# top of the expansion above: about 64 MiB of XML at the defaults, and
# more with a larger payload cap. Plainly timed, opening stored
# element-dense slide XML raised peak memory by 686 MiB for 32 MiB,
# 1,029 MiB for 48 MiB and 1,371 MiB for 64 MiB, about 21 bytes per byte.
# Media count too: real-shaped synthetic decks declared 30 MiB (40 slides
# with 40 photos, near the default 32 MiB payload cap), 32 MiB (an
# embedded 28 MiB video) and 15 MiB (1,000 slides with notes and 200
# pictures), so this allows the default payload cap plus 16 MiB. The
# same figure as the DOCX extractor's.
_MAX_DECLARED_BYTES = 48 * 1024 * 1024

# Members in the package (review round 2). python-pptx builds a part for
# every member a relationship names, about 18 microseconds each plainly
# timed, and tiny stored members add nothing to the expansion above. A
# 1,000-slide deck with notes and media has a few thousand members.
_MAX_MEMBERS = 20_000

# Declared bytes of relationship parts (review round 2). python-pptx reads
# a part's relationships only from the member named for it under
# ``_rels/`` ending ``.rels``, so the name finds every one, and it walks
# each relationship, even to a part that does not exist: 300,000 took
# about 2 s. A relationship is about 100 bytes, so this allows about
# 80,000, several times a large deck's.
_MAX_RELS_BYTES = 8 * 1024 * 1024

# python-pptx 1.0.2 has no constants for the macro-enabled slideshow and
# template main parts (#1042), and its part factory does not map them.
PML_SLIDESHOW_MACRO_MAIN = "application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml"
PML_TEMPLATE_MACRO_MAIN = "application/vnd.ms-powerpoint.template.macroEnabled.main+xml"
PartFactory.part_type_for[PML_SLIDESHOW_MACRO_MAIN] = PresentationPart
PartFactory.part_type_for[PML_TEMPLATE_MACRO_MAIN] = PresentationPart

# Main-part content types this module reads: a presentation, a
# macro-enabled presentation, a slideshow and a template (#947), and a
# macro-enabled slideshow and template (#1042).
_PRESENTATION_MAIN_TYPES = frozenset(
    {
        CT.PML_PRESENTATION_MAIN,
        CT.PML_PRES_MACRO_MAIN,
        CT.PML_SLIDESHOW_MAIN,
        CT.PML_TEMPLATE_MAIN,
        PML_SLIDESHOW_MACRO_MAIN,
        PML_TEMPLATE_MACRO_MAIN,
    }
)

_PARAGRAPH = qn("a:p")
_TEXT_ELEMENTS = (CT_RegularTextRun, CT_TextField, CT_TextLineBreak)


class PptxRelationshipChainError(Exception):
    """The package relates its parts in a chain too long for python-pptx
    to follow: it follows relationships recursively. Fixed text."""

    def __init__(self) -> None:
        super().__init__("package relationship chain too deep to open")


class PptxPackageBudgetError(Exception):
    """The deck's package is over a budget checked before python-pptx
    opens it: expansion, declared size, member count or relationship
    bytes. Fixed text."""

    def __init__(self) -> None:
        super().__init__("package over a PPTX pre-open budget")


def _check_package(payload: bytes) -> None:
    """Raise when the package is over a pre-open budget, read from the
    ZIP central directory (``over_package_budget``)."""
    if over_package_budget(
        payload,
        max_members=_MAX_MEMBERS,
        max_expansion_bytes=_MAX_EXPANSION_BYTES,
        max_rels_bytes=_MAX_RELS_BYTES,
        max_declared_bytes=_MAX_DECLARED_BYTES,
    ):
        raise PptxPackageBudgetError()


def _open_presentation(payload: bytes) -> Presentation:
    """Load a ``.pptx``, ``.pptm``, ``.ppsx``, ``.potx``, ``.ppsm`` or
    ``.potm`` payload.

    The same check ``pptx.Presentation`` makes, widened to the slideshow
    and template types, plain and macro-enabled. Any other main part (a
    document, a workbook) raises ``ValueError`` with fixed text, as
    ``pptx.Presentation`` would.
    """
    part = Package.open(io.BytesIO(payload)).main_document_part
    if part.content_type not in _PRESENTATION_MAIN_TYPES or not isinstance(part, PresentationPart):
        raise ValueError("main part is not a PowerPoint presentation, slideshow or template")
    return part.presentation


class _Budget:
    """The four work budgets of one extraction. ``take`` charges one item
    before it is read; when that item does not fit, the budget it ran out
    of is recorded in ``cut`` and the walk stops."""

    def __init__(self) -> None:
        self.left = {
            "pptx_slides": _MAX_SLIDES,
            "pptx_shapes": _MAX_SHAPES,
            "pptx_table_cells": _MAX_TABLE_CELLS,
            "pptx_text_chars": _MAX_TEXT_CHARS,
        }
        self.cut: str | None = None

    def take(self, cap: str, units: int = 1) -> bool:
        if self.left[cap] < units:
            self.cut = cap
            return False
        self.left[cap] -= units
        return True

    def take_text(self, text: str) -> str:
        """Charge one element and its ``text``, and return the part of the
        text that fits. When not all of it fits, the text budget is spent
        and recorded in ``cut``."""
        left = self.left["pptx_text_chars"]
        cost = _ELEMENT_COST + len(text)
        if cost <= left:
            self.left["pptx_text_chars"] = left - cost
            return text
        self.left["pptx_text_chars"] = 0
        self.cut = "pptx_text_chars"
        return text[: max(left - _ELEMENT_COST, 0)]


def extract(
    payload: bytes,
    *,
    ocr_enabled: bool = True,  # noqa: ARG001
    max_ocr_pages: int = 20,  # noqa: ARG001
    ocr_timeout_seconds: float | None = None,  # noqa: ARG001
    max_pdf_pages: int | None = None,  # noqa: ARG001
    on_progress: Callable[[], None] | None = None,
) -> tuple[str, str]:
    """Extract text from a PPTX, PPTM, PPSX, POTX, PPSM or POTM payload
    in the child process (``ooxml``, #1040). Returns (text, "pptx")."""
    text, caps = run_child(
        "pptx",
        payload,
        max_address_space_bytes=CHILD_MAX_ADDRESS_SPACE_BYTES,
        max_cpu_seconds=CHILD_MAX_CPU_SECONDS,
        timeout_seconds=CHILD_TIMEOUT_SECONDS,
        caps=_CAP_NAMES,
        permanent={"PptxPackageBudgetError": PptxPackageBudgetError},
        on_progress=on_progress,
    )
    for cap in caps:
        warn_extractor_cap(log, cap, "presentation truncated at a walk budget")
    return text, "pptx"


def extract_text(payload: bytes) -> tuple[str, list[str]]:
    """The deck's text and the names of the walk budgets that cut it.
    Runs in the child process (``extractor_child``)."""
    _check_package(payload)
    try:
        presentation = _open_presentation(payload)
    except RecursionError:
        raise PptxRelationshipChainError() from None

    budget = _Budget()
    lines: list[str] = []
    seen: set[str] = set()
    for slide in presentation.slides:
        if not budget.take("pptx_slides"):
            break
        partname = str(slide.part.partname)
        if partname in seen:
            continue
        seen.add(partname)
        if not _slide_lines(slide, budget, lines):
            break

    return "\n\n".join(lines), [budget.cut] if budget.cut is not None else []


def _slide_lines(slide: Slide, budget: _Budget, lines: list[str]) -> bool:
    """Append the slide's shape text, then its notes. False once a budget
    ran out."""
    if not _shape_lines(iter(slide.shapes), budget, lines):
        return False
    if slide.has_notes_slide:
        return _notes_lines(slide.notes_slide, budget, lines)
    return True


def _notes_lines(notes: NotesSlide, budget: _Budget, lines: list[str]) -> bool:
    """Append the text of the notes page's body placeholder, the speaker
    notes. The page's other shapes (slide image, number, header and
    footer) are charged but not read."""
    for shape in notes.shapes:
        if not budget.take("pptx_shapes"):
            return False
        if (
            isinstance(shape, Shape)
            and shape.is_placeholder
            and shape.placeholder_format.type == PP_PLACEHOLDER.BODY
            and not _text_lines(shape._sp.txBody, budget, lines)
        ):
            return False
    return True


def _shape_lines(shapes: Iterator[BaseShape], budget: _Budget, lines: list[str]) -> bool:
    """Append the text of each shape in document order, descending into
    groups with an explicit stack so nesting depth costs no recursion."""
    stack = [shapes]
    while stack:
        shape = next(stack[-1], None)
        if shape is None:
            stack.pop()
            continue
        if not budget.take("pptx_shapes"):
            return False
        if isinstance(shape, GroupShape):
            stack.append(iter(shape.shapes))
        elif isinstance(shape, Shape):
            if not _text_lines(shape._sp.txBody, budget, lines):
                return False
        elif isinstance(shape, GraphicFrame) and shape.has_table:
            if not _table_lines(shape.table, budget, lines):
                return False
    return True


def _table_lines(table: Table, budget: _Budget, lines: list[str]) -> bool:
    """One line per row, its cells' text space-joined; a cell covered by
    a merged cell is skipped, as its text is not shown.

    The rows and each row's cells are listed once, from the table's XML:
    python-pptx's row collection has no ``__iter__``, so iterating it
    indexes it and lists every row again for each row, quadratic in the
    row count (a table of 20,000 empty rows took about two minutes)."""
    for tr in table._tbl.tr_lst:
        if not budget.take("pptx_table_cells"):
            return False
        pieces: list[str] = []
        complete = True
        for tc in tr.tc_lst:
            if not budget.take("pptx_table_cells"):
                complete = False
                break
            if tc.is_spanned:
                continue
            if not _text_lines(tc.txBody, budget, pieces):
                complete = False
                break
        if pieces:
            lines.append(" ".join(pieces))
        if not complete:
            return False
    return True


def _text_lines(tx_body: CT_TextBody | None, budget: _Budget, lines: list[str]) -> bool:
    """Append each non-empty paragraph of a text body. Paragraphs and the
    elements inside them are read one at a time so each is charged as it
    is read; a line break becomes a newline. A shape or cell with no
    text body has nothing to read."""
    if tx_body is None:
        return True
    for paragraph in tx_body.iterchildren(_PARAGRAPH):
        budget.take_text("")
        if budget.cut is not None:
            return False
        pieces: list[str] = []
        for element in paragraph.iterchildren():
            text = element.text if isinstance(element, _TEXT_ELEMENTS) else ""
            pieces.append(budget.take_text(text))
            if budget.cut is not None:
                break
        line = "".join(pieces).replace("\v", "\n").strip()
        if line:
            lines.append(line)
        if budget.cut is not None:
            return False
    return True
