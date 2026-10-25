"""Prefixed OOXML that the existing libraries accept (#1430), synthetic only."""

import io

import pytest
from src import extractors


@pytest.mark.parametrize("module", ["docx", "xlsx"])
@pytest.mark.parametrize(
    "padding", [b"\x00" * 256, b"SYNTHETIC_ZIP_PREFIX\n"], ids=["zero-padding", "text-prefix"]
)
def test_legacy_label_retains_reader_supported_prefixed_ooxml(module, padding):
    marker = "SYNTHETIC_PREFIXED_OFFICE_TEXT"
    stream = io.BytesIO()
    if module == "docx":
        from docx import Document

        document = Document()
        document.core_properties.author = "Synthetic Tester"
        document.core_properties.last_modified_by = "Synthetic Tester"
        document.add_paragraph(marker)
        document.save(stream)
        legacy_type, legacy_name = "application/msword", "synthetic.doc"
        proper_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        cache_module = "doc"
    else:
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.properties.creator = "Synthetic Tester"
        workbook.properties.lastModifiedBy = "Synthetic Tester"
        workbook.active["A1"] = marker
        workbook.save(stream)
        workbook.close()
        legacy_type, legacy_name = "application/vnd.ms-excel", "synthetic.xls"
        proper_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        cache_module = "xls"
    payload = padding + stream.getvalue()
    proper = extractors.extract(
        content_type=proper_type, filename=f"synthetic.{module}", payload=payload
    )
    assert proper.status == extractors.STATUS_SUCCESS
    assert marker in proper.text
    legacy = extractors.extract(content_type=legacy_type, filename=legacy_name, payload=payload)
    assert legacy.status == extractors.STATUS_SUCCESS
    assert legacy.text == proper.text
    assert extractors.extraction_module(legacy_type, legacy_name, payload) == cache_module


@pytest.mark.serial
def test_non_zip_suffix_check_reads_a_fixed_tail_in_the_bounded_child(monkeypatch):
    import time

    from src.extractors import container, container_child

    original = io.BytesIO
    reads = []

    class Counted(original):
        def read(self, size=-1):
            result = super().read(size)
            reads.append(len(result))
            return result

    monkeypatch.setattr(container_child.io, "BytesIO", Counted)
    work = []
    for size in (1024 * 1024, extractors.DEFAULT_MAX_BYTES):
        payload = b"x" * size
        reads.clear()
        start = time.perf_counter()
        assert container_child.extract_text(payload) == (container.NOT_CONTAINER, [])
        assert time.perf_counter() - start < 2
        work.append(sum(reads))
    assert work[0] == work[1]
    assert 0 < work[1] <= 131072


def test_legacy_unstamped_decline_refreshes_once(tmp_path):
    from src.attachment_indexing import apply_attachment_writes, prepare_attachment_writes
    from src.main import _occurrence_reruns_extraction

    from tests.test_attachment_indexing import _attachment, _kwargs, _setup_db_for_attachment

    db = _setup_db_for_attachment(tmp_path)
    attachment = _attachment(
        b"SYNTHETIC_NOT_A_ZIP", filename="synthetic.doc", content_type="application/msword"
    )
    plan = prepare_attachment_writes(db=db, **_kwargs(attachment))
    assert plan.status == extractors.STATUS_UNSUPPORTED
    with db.transaction():
        apply_attachment_writes(plan=plan, claimant_id="msg@x", thread_id="thread-x", db=db)
    db._conn.execute("UPDATE attachment_extractions SET extractor=NULL")
    db._conn.commit()
    assert db.find_no_extractor_attachment_filepaths(_occurrence_reruns_extraction)
    again = prepare_attachment_writes(db=db, **_kwargs(attachment))
    assert again.cached is False
    with db.transaction():
        apply_attachment_writes(plan=again, claimant_id="msg@x", thread_id="thread-x", db=db)
    assert db.find_no_extractor_attachment_filepaths(_occurrence_reruns_extraction) == set()
    assert prepare_attachment_writes(db=db, **_kwargs(attachment)).cached is True
